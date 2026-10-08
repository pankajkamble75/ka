"""plan-04 (research-01 R3, R16 layout extras) — the second extraction pass, spans, row locators, re-extraction, OCR hook.

Characterization cases (P8, N5-before, P9-before) ran green against the pre-plan-04 governance code before Phase 3 changed it.
"""
from __future__ import annotations

from pathlib import Path

from ka.tests.conftest import D
from ka.vocab import AuthorityType

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "underwriting_sop.md"


def _upload_sop(ka, **kw):
    got = ka.ingestion.upload(filename="underwriting_sop.md", data=FIXTURE.read_bytes(), owner="ops", scope=D,
                              authority=AuthorityType.OPERATING_PROCEDURE)
    return got, ka.governance.extract_from_source(got.source, got.version, got.extraction, actor="ops", **kw)


# ---------------------------------------------------------------- Phase 3 characterization (pre-change behaviour)

def test_P8_the_statement_pass_alone_yields_statement_only_candidates_without_subjects(ka):
    """Characterization. Before Phase 3 this ran against unchanged code with no `process_pass` flag and asserted the same
    thing plus 'no span ids' (commit f94b0f3). After Phase 3 the pre-change behaviour is reachable with process_pass=False;
    spans now exist on pass-one evidence too, which is R16's layout extra, so that clause is gone."""
    got, cands = _upload_sop(ka, process_pass=False)
    assert cands and all(c.subject is None and c.predicate is None for c in cands)
    for c in cands:
        ev = ka.repo.evidence.get(c.evidence_refs[0])
        assert ev.locator
    # the statements the modal-verb pass finds (pinned so P9 can compare with the process pass on)
    assert {c.statement for c in cands} == STATEMENTS_BEFORE


STATEMENTS_BEFORE = {
    "Applications with a risk score above 80 must be referred to the credit committee.",
    "Beneficial owner means a natural person who owns 25% or more of the merchant.",
    "Communicate decision — send the outcome to the merchant within two business days.",
    "Make credit decision — approve, decline, or refer the application based on the risk score.",
}


def test_N5_an_empty_source_yields_no_candidates(ka):
    got = ka.ingestion.paste(text="   ", owner="u", scope=D)
    assert ka.governance.extract_from_source(got.source, got.version, got.extraction, actor="u") == []


# ---------------------------------------------------------------- Phase 1: layout extras

import io  # noqa: E402
import json  # noqa: E402

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from ka import config, extraction  # noqa: E402
from ka.api import PREFIX, create_app, set_ka  # noqa: E402
from ka.extraction import EXTRACTION_VERSION, extract_text  # noqa: E402
from ka.llm import StubLLMProvider  # noqa: E402
from ka.process_extraction import AssertionCandidate, ProcessExtractor, render_statement  # noqa: E402
from ka.vocab import ExtractionStatus, SourceType  # noqa: E402


def test_P1_spans_cover_the_stored_text_in_order_and_the_version_is_stamped(ka):
    t = extract_text(FIXTURE.read_bytes(), SourceType.MARKDOWN)
    assert [sp.span_id for sp in t.spans] == ["s1", "s2"] and EXTRACTION_VERSION == "ka-extract/2"
    assert all(t.text[sp.start:sp.end] == t.sections[i][1] for i, sp in enumerate(t.spans))
    assert t.spans[1].start == t.spans[0].end + 2
    got = ka.ingestion.upload(filename="sop.md", data=FIXTURE.read_bytes(), owner="u", scope=D)
    assert got.version.extraction_version == "ka-extract/2"


def test_P2_word_tables_and_sheets_get_row_level_sections():
    import docx
    import openpyxl
    d = docx.Document(); tb = d.add_table(rows=3, cols=2)
    for i, row in enumerate(tb.rows):
        row.cells[0].text, row.cells[1].text = f"k{i}", f"v{i}"
    buf = io.BytesIO(); d.save(buf)
    t = extract_text(buf.getvalue(), SourceType.WORD)
    assert [loc for loc, _ in t.sections] == ["table 1 row 1", "table 1 row 2", "table 1 row 3"] and t.sections[1][1] == "k1 | v1"
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["a", 1]); ws.append(["b", 2])
    buf = io.BytesIO(); wb.save(buf)
    t2 = extract_text(buf.getvalue(), SourceType.SPREADSHEET)
    assert [loc for loc, _ in t2.sections] == ["sheet Sheet row 1", "sheet Sheet row 2"]


def _blank_pdf() -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(width=72, height=72); buf = io.BytesIO(); w.write(buf)
    return buf.getvalue()


def test_P3_ocr_hook_turns_a_textless_pdf_into_page_spans(monkeypatch):
    monkeypatch.setattr(extraction, "ocr_pdf", lambda data: [("p.1", "Refunds above $500 require manager approval.")])
    with config.scoped(KA_OCR=True):
        t = extract_text(_blank_pdf(), SourceType.PDF)
    assert t.status == ExtractionStatus.EXTRACTED and t.note == "ocr" and t.spans[0].span_id == "s1.p1"


def test_N3_without_ocr_a_textless_pdf_stays_unavailable_and_missing_libraries_are_named(monkeypatch):
    with config.scoped(KA_OCR=False):
        t = extract_text(_blank_pdf(), SourceType.PDF)
    assert t.status == ExtractionStatus.UNAVAILABLE and "no text layer" in t.note
    monkeypatch.setattr(extraction, "ocr_pdf", lambda data: None)
    with config.scoped(KA_OCR=True):
        t2 = extract_text(_blank_pdf(), SourceType.PDF)
    assert t2.status == ExtractionStatus.UNAVAILABLE and "ocr" in t2.note and "[ocr]" in t2.note


def test_P4_reextract_creates_a_new_version_with_the_same_checksum_and_leaves_v1_alone(ka):
    got = ka.ingestion.upload(filename="sop.md", data=FIXTURE.read_bytes(), owner="u", scope=D)
    v1 = ka.repo.source_versions.get(got.version.id)
    v1.extraction_version = "ka-extract/1"; ka.repo.source_versions.put(v1)          # pretend an older extractor made it
    again = ka.ingestion.reextract(got.source.id, owner="u")
    assert again.version.version == 2 and again.version.checksum == v1.checksum and again.version.extraction_version == "ka-extract/2"
    assert ka.repo.source_versions.get(v1.id).extraction_version == "ka-extract/1" and ka.repo.sources.get(got.source.id).current_version_id == again.version.id


def test_N6_reextract_without_stored_bytes_is_refused(ka):
    got = ka.ingestion.link(url="http://127.0.0.1:1/blocked", owner="u", scope=D)
    with pytest.raises(ValueError, match="no stored bytes"):
        ka.ingestion.reextract(got.source.id, owner="u")


# ---------------------------------------------------------------- Phase 2: the process pass

def test_P5_heuristic_pass_reads_the_sop_into_description_steps_and_actor():
    t = extract_text(FIXTURE.read_bytes(), SourceType.MARKDOWN)
    got = ProcessExtractor(None, None).extract("SOP", t)
    preds = [a.predicate for a in got.assertions]
    assert got.method == "heuristic" and preds.count("description") == 1 and preds.count("decomposes_into") == 5 and preds.count("performed_by") == 1
    assert [a.object_name for a in got.assertions if a.predicate == "decomposes_into"] == ["Collect application", "Validate application", "Analyze merchant risk", "Make credit decision", "Communicate decision"]
    assert next(a.object_name for a in got.assertions if a.predicate == "performed_by") == "Underwriting team"
    for a in got.assertions:
        sp = t.span(a.span_id)
        assert sp is not None and a.excerpt in t.text[sp.start:sp.end]


def test_N1_the_glossary_section_yields_no_process_assertions():
    t = extract_text(FIXTURE.read_bytes(), SourceType.MARKDOWN)
    got = ProcessExtractor(None, None).extract("SOP", t)
    assert all(a.span_id == "s1" for a in got.assertions) and t.spans[1].locator == "Glossary"


def test_N2_the_heuristic_never_emits_typed_as():
    t = extract_text(b"# Credit review\nThis is a decision process. 1. Read the file. 2. Decide.\n", SourceType.MARKDOWN)
    got = ProcessExtractor(None, None).extract("x", t)
    assert got.assertions and not any(a.predicate == "typed_as" for a in got.assertions)


def test_P6_model_pass_keeps_valid_items_drops_invalid_ones_and_passes_unknown_types_through(ka):
    items = json.dumps([
        {"subject_kind": "process", "subject_name": "Merchant underwriting", "predicate": "typed_as", "object_value": "decision", "excerpt": "Merchant underwriting evaluates", "confidence": 0.8},
        {"subject_kind": "process", "subject_name": "Merchant underwriting", "predicate": "frobnicates", "object_value": "x", "excerpt": "", "confidence": 0.8},
        {"subject_kind": "widget", "subject_name": "Thing", "predicate": "description", "object_value": "x", "excerpt": "", "confidence": 0.8},
        {"subject_kind": "process", "subject_name": "Merchant underwriting", "predicate": "typed_as", "object_value": "frobnication", "excerpt": "not in the text", "confidence": 0.5},
    ])
    t = extract_text(FIXTURE.read_bytes(), SourceType.MARKDOWN)
    got = ProcessExtractor(StubLLMProvider(responses={"PROCESS PASS": items}), ka.grammar).extract("SOP", t)
    assert got.method == "model" and [a.object_value for a in got.assertions if a.predicate == "typed_as"].count("decision") >= 1
    reasons = " ".join(d["reason"] for d in got.dropped)
    assert "frobnicates" in reasons and "widget" in reasons and "excerpt not found" in reasons
    assert any(a.object_value == "frobnication" for a in got.assertions)      # unknown TYPE is kept → binds unresolved


def test_P7_render_statement():
    a = AssertionCandidate("process", "Merchant underwriting", "decomposes_into", "s1", "Merchant underwriting", "x", 0.7, object_kind="process", object_name="Collect application")
    b = AssertionCandidate("process", "Merchant underwriting", "performed_by", "s1", "Merchant underwriting", "x", 0.7, object_kind="actor", object_name="Underwriting team")
    assert render_statement(a) == "Merchant underwriting decomposes into: Collect application"
    assert render_statement(b) == "Merchant underwriting is performed by: Underwriting team"


def test_N4_a_span_the_version_does_not_have_is_dropped_and_a_wrong_excerpt_is_corrected(ka):
    t = extract_text(FIXTURE.read_bytes(), SourceType.MARKDOWN)
    class Fake:
        def extract(self, title, extraction):
            from ka.process_extraction import ProcessExtraction
            good = AssertionCandidate("process", "Merchant underwriting", "description", "s1", "Merchant underwriting", "NOT IN THE TEXT", 0.7, object_value="x", statement="Merchant underwriting: x")
            bad = AssertionCandidate("process", "Merchant underwriting", "description", "s99", "nowhere", "x", 0.7, object_value="y", statement="Merchant underwriting: y")
            return ProcessExtraction(assertions=[good, bad], method="fake")
    ka.governance.process_extractor = Fake()
    got = ka.ingestion.upload(filename="sop.md", data=FIXTURE.read_bytes(), owner="u", scope=D)
    cands = ka.governance.extract_from_source(got.source, got.version, got.extraction, actor="u")
    report = ka.repo.source_versions.get(got.version.id).extraction_report
    assert report["assertions"] == 1 and any("s99" in d["reason"] for d in report["dropped"])
    ev = ka.repo.evidence.get([c for c in cands if c.subject][0].evidence_refs[0])
    assert ev.excerpt in t.text[t.spans[0].start:t.spans[0].end]


# ---------------------------------------------------------------- Phase 3–4: the pipeline and the API

def test_P9_the_statement_pass_is_unchanged_by_the_process_pass(ka):
    got, cands = _upload_sop(ka)
    assert {c.statement for c in cands if c.subject is None} == STATEMENTS_BEFORE


def test_P10_uploading_the_sop_through_the_api_yields_process_candidates_with_spans(ka):
    client = TestClient(create_app(ka))
    try:
        r = client.post(f"{PREFIX}/sources/upload", data={"owner": "ops", "scope_type": "DOMAIN", "scope_id": "merchant-acquiring", "authority": "Operating Procedure"},
                        files={"file": ("underwriting_sop.md", FIXTURE.read_bytes(), "text/markdown")})
        assert r.status_code == 200
        rows = r.json()["candidates"]
        subj = [x for x in rows if x["subject"] == "merchant_underwriting"]
        assert len(subj) == 7 and sum(1 for x in subj if x["predicate"] == "decomposes_into") == 5
        assert r.json()["extraction_report"]["assertions"] == 7 and r.json()["extraction_report"]["dropped"] == []
        src = client.get(f"{PREFIX}/sources/{r.json()['source']['id']}").json()
        assert {e["span_id"] for e in src["evidence"]} <= {"s1", "s2"} and src["versions"][0]["extraction_report"]["statements"] == 4
        sid = r.json()["source"]["id"]
        proc_refs = {x["ref"] for x in subj}
        proc_ev = [e for e in src["evidence"] if any(e["id"] in ka.repo.require_version(ref).evidence_refs for ref in proc_refs)]
        assert proc_ev and all(e["span_id"] == "s1" and e["start"] == 0 for e in proc_ev)
        s = client.get(f"{PREFIX}/subjects/merchant_underwriting").json()
        assert len(s["nuggets"]) == 7 and {n["binding"]["binding_status"] for n in s["nuggets"] if n["predicate"] == "decomposes_into"} == {"bound"}
        # PT1: nothing invented — the SOP states no type, so no typed_as assertion exists
        assert not any(x["predicate"] == "typed_as" for x in subj)
        re_ = client.post(f"{PREFIX}/sources/{sid}/reextract?owner=ops")
        assert re_.status_code == 200 and re_.json()["source_version"]["version"] == 2
    finally:
        set_ka(None)


def test_N8_reextract_404_and_409(ka):
    client = TestClient(create_app(ka))
    try:
        assert client.post(f"{PREFIX}/sources/SRC-nope/reextract").status_code == 404
        got = ka.ingestion.link(url="http://10.0.0.9/x", owner="u", scope=D)
        assert client.post(f"{PREFIX}/sources/{got.source.id}/reextract").status_code == 409
    finally:
        set_ka(None)


def test_PT1_underwriting_sop_yields_an_evidenced_process_profile_without_invention(ka):
    """research-01 PT1 (R2, R3): process subject, description, five nested activities each bound (contains/action) or unresolved,
    every populated field with a span, nothing the SOP does not state."""
    got, cands = _upload_sop(ka)
    procs = [c for c in cands if c.subject and c.subject.canonical_key == "merchant_underwriting"]
    assert any(c.predicate == "description" for c in procs)
    steps = [c for c in procs if c.predicate == "decomposes_into"]
    assert len(steps) == 5
    for c in steps:
        b = ka.repo.binding_for(c.ref)
        assert b.binding_status.value in {"bound", "unresolved"} and (b.edge == "contains" if b.binding_status.value == "bound" else True)
    for c in procs:
        ev = ka.repo.evidence.get(c.evidence_refs[0])
        assert ev.span_id and ev.start is not None and ev.excerpt in got.version.text[ev.start:ev.end]
    assert not any(c.predicate == "typed_as" for c in procs)
