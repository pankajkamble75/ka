"""plan-01 Phase 1 — Knowledge Foundation (§3.1, §4, §5, §6, §7, §14 versioning, §35 search)."""
from __future__ import annotations

import pytest

from ka.extraction import extract_text, normalize, source_type_for_filename
from ka.tests.conftest import A, D, approve_all, ingest_policy
from ka.vocab import ExtractionStatus, NuggetStatus, ScopeType, SourceType


def test_P1_one_document_produces_many_nuggets(ka):
    cands = ingest_policy(ka, D, "KYC is required before activation. Refunds above $500 require manager approval. "
                                 "Chargeback responses must occur within 30 days. Settlement occurs after clearing.")
    assert len(cands) == 4
    assert all(c.status == NuggetStatus.PENDING_REVIEW for c in cands)
    assert {c.canonical_id for c in cands} == {"KN-001", "KN-002", "KN-003", "KN-004"}


def test_P2_raw_source_is_immutable_evidence_with_checksum_and_version(ka):
    got = ka.ingestion.paste(text="Refunds settle within 3 days.", owner="u", scope=D, title="Manual")
    assert got.source.checksum and got.version.version == 1 and got.source.extraction_status == ExtractionStatus.EXTRACTED
    again = ka.ingestion.paste(text="Refunds settle within 3 days.", owner="u", scope=D, title="Manual")
    assert again.source.id != got.source.id or not again.is_new_version   # same bytes never make a new version


def test_P3_relinking_a_changed_url_makes_a_new_source_version(ka):
    first = ka.ingestion.link(url="https://example.test/policy", owner="u", scope=D, prefetched=b"Refunds settle within 3 days.")
    second = ka.ingestion.link(url="https://example.test/policy", owner="u", scope=D, prefetched=b"Refunds settle within 5 days.")
    assert second.source.id == first.source.id and second.is_new_version and second.version.version == 2
    assert ka.repo.source_versions.get(first.version.id).text.endswith("3 days.")   # old version untouched
    assert [e["name"] for e in ka.repo.events()][:1] == ["source.ingested"] and "source.updated" in [e["name"] for e in ka.repo.events()]


def test_P4_source_to_nugget_lineage_is_kept_both_ways(ka):
    cands = ingest_policy(ka, D, "Refunds above $500 require manager approval.")
    v = cands[0]
    assert len(v.source_refs) == 1 and len(v.evidence_refs) == 1
    ev = ka.repo.evidence.get(v.evidence_refs[0])
    assert ev.source_id == v.source_refs[0] and "$500" in ev.excerpt
    assert ka.repo.nuggets.where(lambda n: v.source_refs[0] in n.source_refs)


def test_P5_every_nugget_has_an_explicit_scope_and_hierarchy_fields(ka):
    v = ingest_policy(ka, A, "Merchant A refunds above $1,000 require manager approval.")[0]
    assert v.scope_type == ScopeType.INSTANCE and v.scope_id == "merchant-a"
    assert v.domain_id == "merchant-acquiring" and v.parent_domain_id == "payment-processing" and v.structure_id == "universal"


def test_N1_a_source_without_scope_cannot_be_extracted(ka):
    from ka.governance import GovernanceError
    got = ka.ingestion.paste(text="Refunds settle within 3 days.", owner="u", scope=None)
    with pytest.raises(GovernanceError):
        ka.governance.extract_from_source(got.source, got.version, got.extraction, actor="u")


def test_P6_versions_are_immutable_and_history_stays_queryable(ka):
    from ka.versioning import ImmutableVersionError
    v1 = approve_all(ka, ingest_policy(ka, A, "Merchant A refunds above $500 require manager approval."))[0]
    v2 = ka.governance.propose_revision(v1.canonical_id, statement="Merchant A refunds above $1,000 require manager approval.", by="u", reason="2027 policy")
    approve_all(ka, [v2])
    v1 = ka.repo.require_version(v1.ref)
    assert v1.status == NuggetStatus.SUPERSEDED and v1.superseded_by == v2.ref and v2.supersedes == v1.ref
    assert "$500" in v1.statement and ka.repo.active_version(v1.canonical_id).ref == v2.ref
    v1.statement = "tampered"
    with pytest.raises(ImmutableVersionError):
        ka.versioning.save(v1)


def test_P7_temporal_reconstruction_of_active_knowledge(ka):
    from ka import timeutil
    from datetime import datetime, timezone
    timeutil.freeze(datetime(2025, 1, 1, tzinfo=timezone.utc))
    try:
        v1 = approve_all(ka, ingest_policy(ka, A, "Merchant A refunds above $500 require manager approval."))[0]
        timeutil.freeze(datetime(2025, 6, 1, tzinfo=timezone.utc))
        v2 = ka.governance.propose_revision(v1.canonical_id, statement="Merchant A refunds above $1,000 require manager approval.", by="u", reason="x")
        approve_all(ka, [v2])
    finally:
        timeutil.freeze(None)
    march = [n.ref for n in ka.repo.active_at("2025-03-01T00:00:00+00:00", A)]
    july = [n.ref for n in ka.repo.active_at("2025-07-01T00:00:00+00:00", A)]
    assert march == [v1.ref] and july == [v2.ref]


def test_P8_search_filters_current_pending_and_scope(ka):
    approve_all(ka, ingest_policy(ka, D, "Refunds above $500 require manager approval."))
    ingest_policy(ka, A, "Merchant A requires two signatures on refunds.")
    current = ka.search.search("refunds", filter="Current")
    pending = ka.search.search("refunds", filter="Pending")
    inst = ka.search.search("refunds", filter="Instance")
    assert [h.status for h in current] == ["ACTIVE"]
    assert [h.status for h in pending] == ["PENDING_REVIEW"]
    assert all(h.scope.startswith("INSTANCE") for h in inst)


def test_P9_extractors_cover_markdown_pdf_docx_pptx_xlsx_html():
    import io
    assert source_type_for_filename("a.pdf") == SourceType.PDF and source_type_for_filename("b.docx") == SourceType.WORD
    md = extract_text(b"# Refunds\nRefunds above $500 require approval.", SourceType.MARKDOWN)
    assert md.status == ExtractionStatus.EXTRACTED and md.sections[0][0] == "Refunds"
    import docx
    d = docx.Document(); d.add_heading("Refunds", 1); d.add_paragraph("Refunds above $500 require approval.")
    buf = io.BytesIO(); d.save(buf)
    got = extract_text(buf.getvalue(), SourceType.WORD)
    assert got.status == ExtractionStatus.EXTRACTED and "$500" in got.text and got.sections[0][0] == "Refunds"
    from pptx import Presentation
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "Settlement occurs after clearing."
    buf = io.BytesIO(); prs.save(buf)
    assert "Settlement" in extract_text(buf.getvalue(), SourceType.POWERPOINT).text
    import openpyxl
    wb = openpyxl.Workbook(); wb.active["A1"] = "Threshold"; wb.active["B1"] = 500
    buf = io.BytesIO(); wb.save(buf)
    assert "500" in extract_text(buf.getvalue(), SourceType.SPREADSHEET).text
    html = extract_text(b"<html><title>P</title><body><h2>Refunds</h2><p>Refunds settle within 3 days.</p><script>x()</script></body></html>", SourceType.URL)
    assert "settle" in html.text and "x()" not in html.text
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(width=72, height=72); buf = io.BytesIO(); w.write(buf)
    assert extract_text(buf.getvalue(), SourceType.PDF).status == ExtractionStatus.UNAVAILABLE   # no text layer → evidence kept, not refused


def test_N2_a_broken_file_is_recorded_not_refused(ka):
    got = ka.ingestion.upload(filename="bad.pdf", data=b"not a pdf", owner="u", scope=D)
    assert got.source.extraction_status in {ExtractionStatus.FAILED, ExtractionStatus.UNAVAILABLE}
    assert ka.repo.sources.get(got.source.id) is not None


def test_P10_normalize_folds_currency_commas_and_stopwords():
    assert normalize("Refunds above $1,000 require manager approval.") == normalize("refunds above 1000 require manager approval")
